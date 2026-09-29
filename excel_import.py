from openpyxl import load_workbook
from models import db, Pupil


def import_pupils_from_excel(file_stream, class_id, stream_id=None):
    wb = load_workbook(file_stream, data_only=True)
    ws = wb.active

    header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and any(str(c).strip().upper() == 'NAME' for c in row if c):
            header_row = i
            break
    if not header_row:
        return 0, 0, ['Could not find header row with "NAME"']

    headers = [str(c).strip().upper() if c else '' for c in ws[header_row]]
    name_idx = headers.index('NAME')
    adm_idx = headers.index('ADA') if 'ADA' in headers else None
    phone_idx = headers.index('PHONE') if 'PHONE' in headers else None
    simu_idx = headers.index('SIMU') if 'SIMU' in headers else None

    created, skipped, errors = 0, 0, []

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or not row[name_idx]:
            continue
        raw_name = str(row[name_idx]).strip()
        if not raw_name or raw_name.upper() == 'NAME':
            continue

        if '  ' in raw_name:
            parts = [p.strip() for p in raw_name.split('  ') if p.strip()]
        else:
            parts = raw_name.split()
        if not parts:
            continue

        first_name = parts[0].title()
        last_name = ' '.join(parts[1:]).title() if len(parts) > 1 else ''

        adm = str(row[adm_idx]).strip() if adm_idx is not None and row[adm_idx] else None
        phone = None
        for idx in (phone_idx, simu_idx):
            if idx is not None and row[idx]:
                phone = str(row[idx]).strip()
                break

        existing = None
        if adm:
            existing = Pupil.query.filter_by(admission_no=adm).first()
        if not existing:
            existing = Pupil.query.filter_by(
                first_name=first_name, last_name=last_name, class_id=class_id
            ).first()
        if existing:
            skipped += 1
            continue

        try:
            p = Pupil(first_name=first_name, last_name=last_name,
                      admission_no=adm, class_id=class_id,
                      stream_id=stream_id, parent_phone=phone)
            db.session.add(p)
            db.session.flush()
            created += 1
        except Exception as e:
            db.session.rollback()
            errors.append(f'{raw_name}: {e}')

    db.session.commit()
    return created, skipped, errors